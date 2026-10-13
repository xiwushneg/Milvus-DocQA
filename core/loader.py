"""文档解析模块：PDF / DOCX / Markdown / TXT → 统一段落(segment)列表。

segment 结构: {"text": str, "page": int}
- PDF: 逐页解析（PyMuPDF），page 为真实页码，便于答案溯源
- DOCX: 全文合并为一段（含表格行），page 记为 1
- MD/TXT: 整个文件为一段，page 记为 1
"""
from __future__ import annotations

import hashlib
import io
from pathlib import Path

SUPPORTED_SUFFIXES = {".pdf", ".docx", ".md", ".markdown", ".txt", ".text"}


def file_sha256(data: bytes) -> str:
    """文件内容摘要，用于入库去重。"""
    return hashlib.sha256(data).hexdigest()[:16]


def parse_file(path: str) -> list[dict]:
    p = Path(path)
    if p.suffix.lower() not in SUPPORTED_SUFFIXES:
        raise ValueError(
            f"不支持的文件类型: {p.suffix}（支持 {sorted(SUPPORTED_SUFFIXES)}）"
        )
    return parse_bytes(p.name, p.read_bytes())


def parse_bytes(filename: str, data: bytes) -> list[dict]:
    suffix = Path(filename).suffix.lower()
    if suffix == ".pdf":
        return _parse_pdf(data)
    if suffix == ".docx":
        return _parse_docx(data)
    if suffix in (".md", ".markdown", ".txt", ".text"):
        return [{"text": _decode(data), "page": 1}]
    raise ValueError(f"不支持的文件类型: {suffix}")


def _decode(data: bytes) -> str:
    for enc in ("utf-8", "gb18030", "utf-16"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="ignore")


def _parse_pdf(data: bytes) -> list[dict]:
    import fitz  # PyMuPDF

    segments: list[dict] = []
    with fitz.open(stream=data, filetype="pdf") as doc:
        for i, page in enumerate(doc):
            text = page.get_text("text")
            if text.strip():
                segments.append({"text": text, "page": i + 1})
    return segments


def _parse_docx(data: bytes) -> list[dict]:
    import docx

    d = docx.Document(io.BytesIO(data))
    lines = [p.text for p in d.paragraphs if p.text.strip()]
    for table in d.tables:  # 表格内容一并纳入
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                lines.append(" | ".join(cells))
    return [{"text": "\n".join(lines), "page": 1}]
